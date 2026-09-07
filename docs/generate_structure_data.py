# -*- coding: utf-8 -*-
"""Generate struktur data database (.docx) bergaya tabel skripsi.

Jalankan: python docs/generate_structure_data.py
Output: docs/struktur-data-database.docx
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


# (nama_field, tipe_data, keterangan)
TABLES: list[tuple[str, str, list[tuple[str, str, str]]]] = [
    (
        "users",
        "Tabel Data Users",
        [
            ("id", "bigint(20)", "Primary key, auto-increment, ID unik pengguna"),
            ("name", "varchar(255)", "Nama lengkap pengguna"),
            ("username", "varchar(255)", "Nama pengguna (unik) yang digunakan untuk login"),
            ("email", "varchar(255)", "Alamat email pengguna (unik)"),
            ("email_verified_at", "timestamp", "Tanggal dan waktu email diverifikasi (nullable)"),
            ("password", "varchar(255)", "Password terenkripsi sebagai kredensial login"),
            ("role", "enum('admin','guru','siswa')", "Role / tipe akun untuk otorisasi akses sistem"),
            ("status", "enum('active','inactive')", "Status akun pengguna (aktif / nonaktif)"),
            ("phone", "varchar(20)", "Nomor telepon pengguna (nullable)"),
            ("nisn", "varchar(20)", "NISN siswa (unik, nullable)"),
            ("nip", "varchar(30)", "NIP guru / admin (nullable)"),
            ("class", "varchar(50)", "Kelas siswa (nullable)"),
            ("remember_token", "varchar(100)", "Token untuk fitur remember me (nullable)"),
            ("created_at", "timestamp", "Tanggal dan waktu data dibuat"),
            ("updated_at", "timestamp", "Tanggal dan waktu data diperbarui"),
        ],
    ),
    (
        "equipment",
        "Tabel Data Equipment",
        [
            ("id", "bigint(20)", "Primary key, auto-increment, ID unik inventaris"),
            ("code", "varchar(255)", "Kode inventaris (unik)"),
            ("name", "varchar(255)", "Nama alat atau bahan"),
            ("category", "varchar(255)", "Kategori inventaris"),
            ("item_type", "enum('alat','bahan')", "Jenis inventaris: alat atau bahan"),
            ("stock", "int unsigned", "Jumlah stok total"),
            ("available", "int unsigned", "Jumlah stok yang tersedia untuk dipinjam / diambil"),
            ("qty_baik", "int unsigned", "Jumlah unit dalam kondisi baik"),
            ("qty_rusak_ringan", "int unsigned", "Jumlah unit rusak ringan"),
            ("qty_rusak_berat", "int unsigned", "Jumlah unit rusak berat"),
            ("location", "varchar(255)", "Lokasi penyimpanan (nullable)"),
            ("description", "text", "Deskripsi inventaris (nullable)"),
            ("image_path", "varchar(255)", "Path gambar inventaris (nullable)"),
            ("status", "varchar(20)", "Status inventaris (tersedia / tidak_tersedia)"),
            ("unit", "varchar(255)", "Satuan stok, terutama untuk bahan (nullable)"),
            ("min_stock", "int unsigned", "Batas minimum stok untuk peringatan (nullable)"),
            ("created_at", "timestamp", "Tanggal dan waktu data dibuat"),
            ("updated_at", "timestamp", "Tanggal dan waktu data diperbarui"),
        ],
    ),
    (
        "practicum_schedules",
        "Tabel Data Practicum Schedules",
        [
            ("id", "bigint(20)", "Primary key, auto-increment, ID unik jadwal praktikum"),
            ("code", "varchar(255)", "Kode jadwal (unik)"),
            ("title", "varchar(255)", "Judul jadwal praktikum"),
            ("mata_kuliah", "varchar(255)", "Nama mata pelajaran / mata kuliah"),
            ("jurusan", "varchar(255)", "Jurusan terkait jadwal"),
            ("kelas", "varchar(255)", "Kelas yang mengikuti praktikum"),
            ("tanggal", "date", "Tanggal jadwal khusus (nullable untuk jadwal mingguan)"),
            ("jam_mulai", "time", "Jam mulai praktikum"),
            ("jam_selesai", "time", "Jam selesai praktikum"),
            ("ruangan", "varchar(255)", "Ruangan praktikum (nullable)"),
            ("guru_id", "bigint(20)", "Foreign key ke users (guru pengampu)"),
            ("priority", "enum('normal','tinggi','lomba')", "Prioritas jadwal praktikum"),
            ("type", "enum('mingguan','khusus')", "Tipe jadwal: berulang mingguan atau khusus"),
            ("hari", "enum('senin','selasa','rabu','kamis','jumat','sabtu','minggu')", "Hari berulang untuk jadwal mingguan (nullable)"),
            ("notes", "text", "Catatan tambahan jadwal (nullable)"),
            ("created_at", "timestamp", "Tanggal dan waktu data dibuat"),
            ("updated_at", "timestamp", "Tanggal dan waktu data diperbarui"),
        ],
    ),
    (
        "submissions",
        "Tabel Data Submissions",
        [
            ("id", "bigint(20)", "Primary key, auto-increment, ID unik pengajuan"),
            ("code", "varchar(255)", "Kode pengajuan (unik)"),
            ("borrower_id", "bigint(20)", "Foreign key ke users (peminjam / siswa)"),
            ("supervisor_id", "bigint(20)", "Foreign key ke users (guru pembimbing, nullable)"),
            ("purpose", "varchar(255)", "Keperluan / tujuan pengajuan"),
            ("notes", "text", "Catatan tambahan pengajuan (nullable)"),
            ("request_date", "date", "Tanggal pengajuan dibuat"),
            ("created_at", "timestamp", "Tanggal dan waktu data dibuat"),
            ("updated_at", "timestamp", "Tanggal dan waktu data diperbarui"),
        ],
    ),
    (
        "loans",
        "Tabel Data Loans",
        [
            ("id", "bigint(20)", "Primary key, auto-increment, ID unik peminjaman / pengambilan"),
            ("code", "varchar(255)", "Kode transaksi (unik)"),
            ("loan_group_id", "char(36)", "UUID pengelompokan transaksi lama (nullable)"),
            ("submission_id", "bigint(20)", "Foreign key ke submissions (nullable)"),
            ("borrower_id", "bigint(20)", "Foreign key ke users (peminjam)"),
            ("supervisor_id", "bigint(20)", "Foreign key ke users (guru pembimbing, nullable)"),
            ("practicum_schedule_id", "bigint(20)", "Foreign key ke practicum_schedules (nullable)"),
            ("item_type", "enum('alat','bahan')", "Jenis transaksi: peminjaman alat atau pengambilan bahan"),
            (
                "status",
                "enum('diminta','antrian','disetujui','ditolak','dipinjam','terlambat','menunggu_inspeksi','dikembalikan','dibatalkan')",
                "Status alur peminjaman / pengambilan",
            ),
            ("queue_priority", "smallint unsigned", "Nilai prioritas antrean (semakin tinggi semakin diprioritaskan)"),
            ("queued_at", "timestamp", "Waktu masuk antrean (nullable)"),
            ("queue_priority_note", "varchar(255)", "Catatan penyesuaian prioritas antrean (nullable)"),
            ("queue_priority_set_by", "bigint(20)", "Foreign key ke users (admin yang mengatur prioritas, nullable)"),
            ("queue_priority_set_at", "timestamp", "Waktu prioritas antrean diatur (nullable)"),
            ("request_date", "date", "Tanggal permintaan"),
            ("borrowed_at", "datetime", "Waktu barang dipinjam / diambil (nullable)"),
            ("due_at", "datetime", "Batas waktu pengembalian (nullable)"),
            ("returned_at", "datetime", "Waktu pengembalian selesai (nullable)"),
            ("purpose", "varchar(255)", "Keperluan transaksi (nullable)"),
            ("notes", "text", "Catatan tambahan (nullable)"),
            ("rejection_reason", "text", "Alasan penolakan (nullable)"),
            ("borrow_scope", "enum('lab','bawa_pulang')", "Lingkup pemakaian: di lab atau bawa pulang"),
            ("borrow_reason", "enum('reguler','lanjutan')", "Alasan peminjaman (nullable)"),
            ("usage_room", "varchar(100)", "Ruangan penggunaan saat di lab (nullable)"),
            ("created_at", "timestamp", "Tanggal dan waktu data dibuat"),
            ("updated_at", "timestamp", "Tanggal dan waktu data diperbarui"),
        ],
    ),
    (
        "loan_items",
        "Tabel Data Loan Items",
        [
            ("id", "bigint(20)", "Primary key, auto-increment, ID unik item transaksi"),
            ("loan_id", "bigint(20)", "Foreign key ke loans"),
            ("equipment_id", "bigint(20)", "Foreign key ke equipment"),
            ("quantity", "int unsigned", "Jumlah item yang diajukan / dipinjam"),
            ("created_at", "timestamp", "Tanggal dan waktu data dibuat"),
            ("updated_at", "timestamp", "Tanggal dan waktu data diperbarui"),
        ],
    ),
    (
        "loan_status_logs",
        "Tabel Data Loan Status Logs",
        [
            ("id", "bigint(20)", "Primary key, auto-increment, ID unik log status"),
            ("loan_id", "bigint(20)", "Foreign key ke loans"),
            ("status", "varchar(255)", "Status yang dicatat pada log"),
            ("note", "text", "Catatan perubahan status (nullable)"),
            ("user_id", "bigint(20)", "Foreign key ke users (pelaku perubahan, nullable)"),
            ("created_at", "timestamp", "Tanggal dan waktu log dibuat"),
        ],
    ),
    (
        "loan_collaterals",
        "Tabel Data Loan Collaterals",
        [
            ("id", "bigint(20)", "Primary key, auto-increment, ID unik jaminan kartu"),
            ("code", "varchar(255)", "Kode jaminan (unik)"),
            ("loan_id", "bigint(20)", "Foreign key ke loans"),
            ("student_id", "bigint(20)", "Foreign key ke users (siswa pemilik kartu)"),
            ("card_type", "enum('kartu_pelajar','kartu_siswa','lainnya')", "Jenis kartu jaminan"),
            ("card_number", "varchar(255)", "Nomor / identitas kartu (nullable)"),
            (
                "status",
                "enum('dititipkan','ditahan','menunggu_kompensasi','dikembalikan','dibatalkan')",
                "Status jaminan kartu",
            ),
            ("held_at", "datetime", "Waktu kartu mulai ditahan (nullable)"),
            ("returned_at", "datetime", "Waktu kartu dikembalikan (nullable)"),
            ("held_by_admin_id", "bigint(20)", "Foreign key ke users (admin yang menahan, nullable)"),
            ("notes", "text", "Catatan jaminan (nullable)"),
            ("created_at", "timestamp", "Tanggal dan waktu data dibuat"),
            ("updated_at", "timestamp", "Tanggal dan waktu data diperbarui"),
        ],
    ),
    (
        "loan_return_inspections",
        "Tabel Data Loan Return Inspections",
        [
            ("id", "bigint(20)", "Primary key, auto-increment, ID unik inspeksi pengembalian"),
            ("loan_id", "bigint(20)", "Foreign key ke loans"),
            ("result", "enum('belum','lengkap','tidak_lengkap','rusak')", "Hasil inspeksi pengembalian"),
            ("notes", "text", "Catatan hasil inspeksi (nullable)"),
            ("missing_items", "text", "Daftar item yang kurang (nullable)"),
            ("damage_description", "text", "Deskripsi kerusakan (nullable)"),
            ("checked_by_admin_id", "bigint(20)", "Foreign key ke users (admin pemeriksa, nullable)"),
            ("checked_at", "timestamp", "Waktu inspeksi dilakukan (nullable)"),
            ("created_at", "timestamp", "Tanggal dan waktu data dibuat"),
            ("updated_at", "timestamp", "Tanggal dan waktu data diperbarui"),
        ],
    ),
    (
        "loan_compensations",
        "Tabel Data Loan Compensations",
        [
            ("id", "bigint(20)", "Primary key, auto-increment, ID unik kompensasi"),
            ("loan_id", "bigint(20)", "Foreign key ke loans"),
            ("required", "tinyint(1)", "Menandai apakah kompensasi wajib"),
            ("status", "enum('tidak_perlu','pending','selesai')", "Status proses kompensasi"),
            ("amount", "int unsigned", "Nilai kompensasi (nullable)"),
            ("description", "text", "Deskripsi kompensasi (nullable)"),
            ("completed_at", "timestamp", "Waktu kompensasi diselesaikan (nullable)"),
            ("completed_by_admin_id", "bigint(20)", "Foreign key ke users (admin yang menyelesaikan, nullable)"),
            ("created_at", "timestamp", "Tanggal dan waktu data dibuat"),
            ("updated_at", "timestamp", "Tanggal dan waktu data diperbarui"),
        ],
    ),
    (
        "notifications",
        "Tabel Data Notifications",
        [
            ("id", "char(36)", "Primary key UUID notifikasi"),
            ("type", "varchar(255)", "Kelas / tipe notifikasi"),
            ("notifiable_type", "varchar(255)", "Tipe model penerima (polymorphic)"),
            ("notifiable_id", "bigint(20)", "ID model penerima (polymorphic)"),
            ("data", "text", "Payload data notifikasi (JSON)"),
            ("read_at", "timestamp", "Waktu notifikasi dibaca (nullable)"),
            ("created_at", "timestamp", "Tanggal dan waktu data dibuat"),
            ("updated_at", "timestamp", "Tanggal dan waktu data diperbarui"),
        ],
    ),
    (
        "push_subscriptions",
        "Tabel Data Push Subscriptions",
        [
            ("id", "bigint(20)", "Primary key, auto-increment, ID unik langganan push"),
            ("subscribable_type", "varchar(255)", "Tipe model pelanggan (polymorphic)"),
            ("subscribable_id", "bigint(20)", "ID model pelanggan (polymorphic)"),
            ("endpoint", "varchar(500)", "Endpoint Web Push (unik)"),
            ("public_key", "varchar(255)", "Public key langganan (nullable)"),
            ("auth_token", "varchar(255)", "Auth token langganan (nullable)"),
            ("content_encoding", "varchar(255)", "Content encoding Web Push (nullable)"),
            ("created_at", "timestamp", "Tanggal dan waktu data dibuat"),
            ("updated_at", "timestamp", "Tanggal dan waktu data diperbarui"),
        ],
    ),
    (
        "password_reset_tokens",
        "Tabel Data Password Reset Tokens",
        [
            ("email", "varchar(255)", "Primary key, email pengguna yang meminta reset"),
            ("token", "varchar(255)", "Token reset password"),
            ("created_at", "timestamp", "Tanggal dan waktu token dibuat (nullable)"),
        ],
    ),
    (
        "sessions",
        "Tabel Data Sessions",
        [
            ("id", "varchar(255)", "Primary key, ID sesi"),
            ("user_id", "bigint(20)", "Foreign key ke users (nullable)"),
            ("ip_address", "varchar(45)", "Alamat IP pengguna (nullable)"),
            ("user_agent", "text", "User agent browser (nullable)"),
            ("payload", "longtext", "Data payload sesi"),
            ("last_activity", "int", "Unix timestamp aktivitas terakhir"),
        ],
    ),
]


HEADER_BG = "BDD7EE"  # biru muda mirip contoh


def set_run_font(run, *, bold: bool = False, size: int = 11, name: str = "Times New Roman"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def shade_cell(cell, hex_color: str):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def set_cell_borders(cell):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "000000")
        tc_borders.append(element)
    tc_pr.append(tc_borders)


def fill_cell(cell, text: str, *, bold: bool = False, center: bool = False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    )
    run = paragraph.add_run(text)
    set_run_font(run, bold=bold, size=11)
    set_cell_borders(cell)


def add_structure_table(
    doc: Document,
    number: float,
    title: str,
    fields: list[tuple[str, str, str]],
):
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(12)
    caption.paragraph_format.space_after = Pt(6)
    run = caption.add_run(f"Tabel 3.{number:g} {title}")
    set_run_font(run, bold=True, size=12)

    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    headers = ["No", "Nama Field", "Tipe Data", "Keterangan"]
    widths = [Cm(1.2), Cm(4.2), Cm(5.5), Cm(6.5)]

    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        fill_cell(cell, header, bold=True, center=True)
        shade_cell(cell, HEADER_BG)
        cell.width = widths[idx]

    for i, (name, tipe, ket) in enumerate(fields, start=1):
        row = table.add_row().cells
        fill_cell(row[0], str(i), center=True)
        fill_cell(row[1], name)
        fill_cell(row[2], tipe)
        fill_cell(row[3], ket)
        for col, width in enumerate(widths):
            row[col].width = width


def build_document() -> Document:
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run("3.x Struktur Basis Data")
    set_run_font(run, bold=True, size=14)

    intro = doc.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    intro_run = intro.add_run(
        "Struktur basis data pada sistem manajemen peminjaman alat dan bahan "
        "Laboratorium Teknik Audio Video digambarkan melalui spesifikasi tabel "
        "berikut. Setiap tabel memuat nama field, tipe data, serta keterangan "
        "fungsi kolom sesuai skema database aplikasi."
    )
    set_run_font(intro_run, size=12)

    start_number = 1
    for offset, (_table_name, title, fields) in enumerate(TABLES):
        add_structure_table(doc, start_number + offset, title, fields)

    return doc


def main():
    out = Path(__file__).resolve().parent / "struktur-data-database.docx"
    doc = build_document()
    doc.save(out)
    print(f"Saved: {out}")


if __name__ == "__main__":
    main()
