# -*- coding: utf-8 -*-
"""Generate dokumen Pengujian Alpha (Black-Box) untuk Lab Audio Video.

Jalankan: python docs/generate_alpha_testing.py
Output: docs/pengujian-alpha-blackbox.docx
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


HEADER_BG = "BDD7EE"

# Setiap section: (heading_letter_label, table_caption, rows[(test_case, expected)])
# Nomor tabel mulai dari 3.21 mengikuti contoh skripsi.
SECTIONS: list[tuple[str, str, list[tuple[str, str]]]] = [
    (
        "a. Pengujian alpha login Admin, Guru, dan Siswa",
        "Pengujian alpha login",
        [
            (
                "Menekan tombol login dengan username dan password benar sebagai Admin (contoh: admin / password)",
                "Sistem menampilkan Dashboard Admin sesuai hak akses role admin",
            ),
            (
                "Menekan tombol login dengan username dan password benar sebagai Guru",
                "Sistem menampilkan Dashboard Guru sesuai hak akses role guru",
            ),
            (
                "Menekan tombol login dengan username dan password benar sebagai Siswa",
                "Sistem menampilkan Dashboard Siswa sesuai hak akses role siswa",
            ),
            (
                "Menekan tombol login dengan username atau password salah",
                "Sistem menampilkan pesan kesalahan kredensial (These credentials do not match our records.) pada field username",
            ),
            (
                "Menekan tombol login tanpa mengisi username dan/atau password",
                "Sistem mencegah submit (validasi wajib isi pada form login)",
            ),
            (
                "Melakukan logout dari akun yang sedang aktif",
                "Sistem mengakhiri sesi dan mengarahkan ke halaman login",
            ),
        ],
    ),
    (
        "b. Pengujian alpha otorisasi antar role (RBAC)",
        "Pengujian alpha otorisasi role",
        [
            (
                "Pengguna role Siswa mengakses URL Admin (contoh: /admin/dashboard)",
                "Sistem menolak akses (HTTP 403 Unauthorized)",
            ),
            (
                "Pengguna role Guru mengakses URL Admin (contoh: /admin/users)",
                "Sistem menolak akses (HTTP 403 Unauthorized)",
            ),
            (
                "Pengguna role Admin mengakses URL Siswa (contoh: /siswa/loans)",
                "Sistem menolak akses (HTTP 403 Unauthorized)",
            ),
            (
                "Pengguna role Guru membuka detail peminjaman yang bukan bimbingan sendiri",
                "Sistem menolak atau tidak menampilkan data di luar supervision Guru",
            ),
            (
                "Pengguna role Siswa membuka detail pengajuan milik siswa lain",
                "Sistem menolak akses terhadap data milik pengguna lain",
            ),
        ],
    ),
    (
        "c. Pengujian alpha kelola pengguna (Admin)",
        "Pengujian alpha kelola pengguna",
        [
            (
                "Admin membuka menu Kelola Pengguna",
                "Sistem menampilkan daftar pengguna beserta filter pencarian dan filter role",
            ),
            (
                "Admin menambahkan pengguna baru role Siswa lengkap dengan kelas dan NISN",
                "Sistem menyimpan data dan menampilkan pesan “Pengguna berhasil ditambahkan.”",
            ),
            (
                "Admin menambahkan pengguna baru role Guru lengkap dengan NIP",
                "Sistem menyimpan data dan menampilkan pesan “Pengguna berhasil ditambahkan.”",
            ),
            (
                "Admin mengubah data pengguna yang sudah ada",
                "Sistem memperbarui data dan menampilkan pesan “Pengguna berhasil diperbarui.”",
            ),
            (
                "Admin mereset kata sandi pengguna",
                "Sistem mereset kata sandi dan menampilkan pesan “Kata sandi berhasil direset.”",
            ),
            (
                "Admin menghapus pengguna",
                "Sistem menghapus data dan menampilkan pesan “Pengguna berhasil dihapus.”",
            ),
            (
                "Admin mencari pengguna berdasarkan nama / username / email",
                "Sistem menampilkan hasil pencarian sesuai kata kunci",
            ),
            (
                "Admin memfilter daftar pengguna berdasarkan role (admin / guru / siswa)",
                "Sistem hanya menampilkan pengguna sesuai role yang dipilih",
            ),
        ],
    ),
    (
        "d. Pengujian alpha kelola alat (Admin)",
        "Pengujian alpha kelola alat",
        [
            (
                "Admin membuka menu Kelola Alat",
                "Sistem menampilkan daftar inventaris alat beserta filter pencarian, kategori, status, dan kondisi",
            ),
            (
                "Admin menambahkan alat baru dengan data lengkap (kode, nama, stok, kondisi, status tersedia)",
                "Sistem menyimpan data dan menampilkan pesan “Alat berhasil ditambahkan.”",
            ),
            (
                "Admin menambahkan / mengubah alat dengan jumlah kondisi (baik + rusak ringan + rusak berat) tidak sama dengan total stok",
                "Sistem menolak penyimpanan dan menampilkan validasi jumlah kondisi harus sama dengan total stok",
            ),
            (
                "Admin mengubah data alat yang sudah ada",
                "Sistem memperbarui data dan menampilkan pesan “Alat berhasil diperbarui.”",
            ),
            (
                "Admin mengubah status alat menjadi tidak tersedia",
                "Sistem memperbarui status inventaris menjadi tidak_tersedia",
            ),
            (
                "Admin menghapus alat",
                "Sistem menghapus data dan menampilkan pesan “Alat berhasil dihapus.”",
            ),
            (
                "Admin mencari alat berdasarkan nama / kode",
                "Sistem menampilkan hasil pencarian sesuai kata kunci",
            ),
        ],
    ),
    (
        "e. Pengujian alpha kelola bahan (Admin)",
        "Pengujian alpha kelola bahan",
        [
            (
                "Admin membuka menu Kelola Bahan",
                "Sistem menampilkan daftar bahan beserta filter pencarian, kategori, dan status",
            ),
            (
                "Admin menambahkan bahan baru dengan stok, satuan, dan status tersedia",
                "Sistem menyimpan data dan menampilkan pesan “Bahan berhasil ditambahkan.”",
            ),
            (
                "Admin mengubah data bahan yang sudah ada",
                "Sistem memperbarui data dan menampilkan pesan “Bahan berhasil diperbarui.”",
            ),
            (
                "Admin menghapus bahan",
                "Sistem menghapus data dan menampilkan pesan “Bahan berhasil dihapus.”",
            ),
            (
                "Admin mencari bahan berdasarkan nama / kode",
                "Sistem menampilkan hasil pencarian sesuai kata kunci",
            ),
        ],
    ),
    (
        "f. Pengujian alpha jadwal praktikum (Admin)",
        "Pengujian alpha jadwal praktikum",
        [
            (
                "Admin membuka menu Jadwal Praktikum",
                "Sistem menampilkan ringkasan minggu ini dan daftar jadwal lengkap",
            ),
            (
                "Admin menambahkan jadwal mingguan (hari, jam, kelas, guru, mata pelajaran)",
                "Sistem menyimpan data dan menampilkan pesan “Jadwal praktikum berhasil ditambahkan.”",
            ),
            (
                "Admin menambahkan jadwal khusus / acara khusus dengan tanggal tertentu",
                "Sistem menyimpan jadwal bertipe khusus sesuai tanggal yang diisi",
            ),
            (
                "Admin mengubah data jadwal praktikum",
                "Sistem memperbarui data dan menampilkan pesan “Jadwal praktikum berhasil diperbarui.”",
            ),
            (
                "Admin menghapus jadwal praktikum",
                "Sistem menghapus data dan menampilkan pesan “Jadwal praktikum berhasil dihapus.”",
            ),
            (
                "Guru membuka menu Jadwal Praktikum",
                "Sistem menampilkan daftar / detail jadwal praktikum (tanpa tombol tambah/ubah/hapus)",
            ),
        ],
    ),
    (
        "g. Pengujian alpha katalog alat dan bahan (Siswa & Guru)",
        "Pengujian alpha katalog inventaris",
        [
            (
                "Siswa membuka menu Alat Lab",
                "Sistem menampilkan tabel katalog alat yang tersedia untuk dilihat",
            ),
            (
                "Siswa membuka detail salah satu alat",
                "Sistem menampilkan informasi alat (stok, status, deskripsi) tanpa tombol ubah/hapus",
            ),
            (
                "Siswa membuka menu Bahan Lab",
                "Sistem menampilkan tabel katalog bahan",
            ),
            (
                "Siswa membuka detail salah satu bahan",
                "Sistem menampilkan informasi bahan (stok, satuan, status) tanpa tombol ubah/hapus",
            ),
            (
                "Guru membuka menu Inventaris (tab alat / bahan)",
                "Sistem menampilkan inventaris lab bersifat baca saja",
            ),
        ],
    ),
    (
        "h. Pengujian alpha pengajuan alat / bahan (Siswa)",
        "Pengujian alpha pengajuan siswa",
        [
            (
                "Siswa membuka halaman Ajukan Alat / Bahan",
                "Sistem menampilkan katalog alat/bahan, keranjang, dan form pengajuan",
            ),
            (
                "Siswa menambahkan alat ke keranjang lalu mengirim pengajuan pakai di lab dengan data lengkap (jadwal hari ini, guru, ruang)",
                "Sistem membuat pengajuan (kode SUB-…) berstatus menunggu persetujuan / antrian dan menampilkan pesan sukses pengajuan",
            ),
            (
                "Siswa mengirim pengajuan alat bawa pulang tanpa menyetujui pernyataan jaminan kartu",
                "Sistem menolak submit dan menampilkan validasi bahwa jaminan kartu pelajar wajib dipahami / disetujui",
            ),
            (
                "Siswa mengirim pengajuan alat bawa pulang dengan pernyataan jaminan disetujui",
                "Sistem membuat pengajuan alat beserta data jaminan kartu (status belum diterima / dititipkan) dan menampilkan pesan sukses",
            ),
            (
                "Siswa menambahkan bahan ke keranjang lalu mengirim pengajuan pengambilan bahan",
                "Sistem membuat pengajuan bahan dan menampilkan pesan sukses pengajuan",
            ),
            (
                "Siswa mengajukan alat dan bahan sekaligus dalam satu keranjang (package)",
                "Sistem membuat satu pengajuan SUB dengan proses alat dan bahan terpisah, menampilkan pesan bahwa alat dan bahan diproses sesuai ketersediaan stok",
            ),
            (
                "Siswa mengajukan alat yang stoknya tidak mencukupi",
                "Sistem memasukkan pengajuan ke status Antrian (Round Robin) atau menampilkan pesan terkait antrean / stok",
            ),
            (
                "Siswa mengirim pengajuan pakai di lab tanpa memilih jadwal / ruang yang wajib",
                "Sistem menolak submit dan menampilkan validasi field wajib (jadwal / ruang / guru sesuai aturan)",
            ),
            (
                "Siswa mengubah pengajuan yang masih berstatus diminta / antrian / disetujui",
                "Sistem menyimpan perubahan dan menampilkan pesan “Pengajuan peminjaman berhasil diperbarui.”",
            ),
            (
                "Siswa membatalkan pengajuan yang masih berstatus diminta / antrian / disetujui",
                "Sistem membatalkan pengajuan dan menampilkan pesan “Pengajuan peminjaman dibatalkan.”",
            ),
        ],
    ),
    (
        "i. Pengujian alpha verifikasi peminjaman (Admin)",
        "Pengujian alpha verifikasi peminjaman",
        [
            (
                "Admin membuka menu Peminjaman",
                "Sistem menampilkan daftar pengajuan siswa beserta filter status, jenis, peminjam, guru, kelas, dan tanggal",
            ),
            (
                "Admin membuka detail pengajuan (SUB-…)",
                "Sistem menampilkan ringkasan pengajuan beserta detail pinjaman alat dan/atau bahan",
            ),
            (
                "Admin menekan Setujui pada pengajuan alat berstatus diminta dengan stok mencukupi",
                "Status menjadi Disetujui, stok terreservasi, dan sistem menampilkan pesan “Peminjaman berhasil disetujui.”",
            ),
            (
                "Admin menekan Setujui pada pengajuan bahan berstatus diminta dengan stok mencukupi",
                "Status bahan menjadi Diambil (dipinjam), stok berkurang, dan sistem menampilkan pesan sukses persetujuan",
            ),
            (
                "Admin menekan Setujui pada pengajuan yang stoknya tidak mencukupi",
                "Sistem tetap/mengembalikan ke Antrian stok dan menampilkan pesan bahwa pengajuan masih dalam antrian stok",
            ),
            (
                "Admin menekan Tolak pada pengajuan dengan mengisi alasan penolakan",
                "Status menjadi Ditolak dan sistem menampilkan pesan “Peminjaman ditolak.”",
            ),
            (
                "Admin menekan Tolak tanpa mengisi alasan",
                "Sistem menolak aksi dan meminta alasan penolakan",
            ),
            (
                "Admin menekan Tandai Dipinjam pada alat berstatus disetujui (pakai di lab)",
                "Status menjadi Dipinjam, due_at terisi, dan sistem menampilkan pesan status alat sedang dipinjam",
            ),
            (
                "Admin mencoba Tandai Dipinjam pada alat bawa pulang sebelum kartu diterima",
                "Sistem menolak aksi dengan pesan bahwa kartu pelajar harus diterima terlebih dahulu",
            ),
            (
                "Admin mengatur prioritas antrian pada pengajuan berstatus antrian",
                "Sistem memperbarui prioritas dan menampilkan pesan “Prioritas antrian berhasil diperbarui.”",
            ),
            (
                "Admin mereset prioritas antrian ke Round Robin (FIFO)",
                "Sistem mereset prioritas dan menampilkan pesan prioritas direset ke Round Robin (FIFO)",
            ),
            (
                "Admin memfilter daftar peminjaman berdasarkan status / jenis / kelas",
                "Sistem menampilkan data sesuai filter aktif",
            ),
        ],
    ),
    (
        "j. Pengujian alpha jaminan kartu (Admin)",
        "Pengujian alpha jaminan kartu",
        [
            (
                "Admin membuka menu Jaminan Kartu",
                "Sistem menampilkan daftar jaminan beserta filter status, siswa, kelas, dan tanggal",
            ),
            (
                "Admin menekan Terima Kartu Pelajar pada jaminan berstatus belum diterima (dititipkan)",
                "Status menjadi Ditahan dan sistem menampilkan pesan “Kartu pelajar diterima sebagai jaminan.”",
            ),
            (
                "Admin mencoba menerima kartu yang sudah ditahan",
                "Sistem menolak aksi karena kartu hanya dapat diterima dari status belum diterima",
            ),
            (
                "Admin menekan Kembalikan Kartu setelah peminjaman alat selesai dan inspeksi lengkap",
                "Status menjadi Sudah dikembalikan dan sistem menampilkan pesan kartu berhasil dikembalikan ke siswa",
            ),
            (
                "Admin mencoba mengembalikan kartu sebelum alat selesai dikembalikan",
                "Sistem menolak aksi dengan pesan bahwa alat harus dikembalikan terlebih dahulu",
            ),
            (
                "Admin menambahkan / mengubah / menghapus data jaminan secara manual",
                "Sistem menampilkan pesan sukses sesuai aksi (ditambahkan / diperbarui / dihapus)",
            ),
        ],
    ),
    (
        "k. Pengujian alpha pengembalian, inspeksi, dan kompensasi",
        "Pengujian alpha pengembalian dan inspeksi",
        [
            (
                "Siswa menekan Ajukan Pengembalian pada peminjaman alat berstatus dipinjam / terlambat",
                "Status menjadi Menunggu Inspeksi dan sistem menampilkan pesan pengembalian diajukan menunggu inspeksi admin",
            ),
            (
                "Admin menekan Ajukan Inspeksi dari sisi admin pada alat dipinjam / terlambat",
                "Status menjadi Menunggu Inspeksi dan sistem menampilkan pesan pengembalian diajukan",
            ),
            (
                "Admin menyelesaikan inspeksi dengan hasil lengkap",
                "Status pinjaman menjadi Dikembalikan, stok alat dipulihkan, dan sistem menampilkan pesan “Inspeksi pengembalian selesai.”",
            ),
            (
                "Admin menyelesaikan inspeksi dengan hasil tidak lengkap tanpa mengisi item yang kurang",
                "Sistem menolak aksi dan menampilkan validasi “Item yang kurang wajib diisi.”",
            ),
            (
                "Admin menyelesaikan inspeksi dengan hasil rusak tanpa mengisi deskripsi kerusakan",
                "Sistem menolak aksi dan menampilkan validasi deskripsi / tingkat kerusakan wajib diisi",
            ),
            (
                "Admin menyelesaikan inspeksi dengan hasil tidak lengkap / rusak",
                "Status pinjaman selesai (dikembalikan), jaminan masuk Menunggu Kompensasi, dan kompensasi berstatus pending",
            ),
            (
                "Admin menandai Kompensasi Selesai pada jaminan menunggu kompensasi",
                "Kompensasi selesai, kartu dikembalikan ke siswa, dan sistem menampilkan pesan “Kompensasi selesai. Kartu dikembalikan.”",
            ),
            (
                "Siswa membuka dashboard saat masih ada kompensasi pending",
                "Sistem menampilkan peringatan kompensasi / alat rusak pada dashboard siswa",
            ),
        ],
    ),
    (
        "l. Pengujian alpha daftar & riwayat peminjaman (Siswa & Guru)",
        "Pengujian alpha daftar dan riwayat",
        [
            (
                "Siswa membuka menu Alat & Bahan Saya",
                "Sistem menampilkan daftar pengajuan aktif beserta status agregat",
            ),
            (
                "Siswa membuka menu Riwayat Alat & Bahan",
                "Sistem menampilkan arsip pengajuan selesai / dibatalkan",
            ),
            (
                "Siswa mencari pengajuan berdasarkan kode / barang / keperluan",
                "Sistem menampilkan hasil sesuai kata kunci",
            ),
            (
                "Siswa memfilter daftar berdasarkan status dan rentang tanggal",
                "Sistem menampilkan data sesuai filter",
            ),
            (
                "Guru membuka menu Peminjaman Siswa",
                "Sistem menampilkan peminjaman aktif siswa bimbingan Guru",
            ),
            (
                "Guru membuka Riwayat Peminjaman Siswa (scope history)",
                "Sistem menampilkan arsip peminjaman bimbingan Guru",
            ),
            (
                "Guru membuka detail pengajuan / pinjaman bimbingan",
                "Sistem menampilkan detail status, item, dan timeline tanpa aksi verifikasi admin",
            ),
        ],
    ),
    (
        "m. Pengujian alpha dashboard tiap role",
        "Pengujian alpha dashboard",
        [
            (
                "Admin membuka Dashboard Admin",
                "Sistem menampilkan ringkasan operasional (permintaan pending, antrian, keterlambatan, stok menipis, dll.)",
            ),
            (
                "Guru membuka Dashboard Guru",
                "Sistem menampilkan ringkasan siswa pinjam aktif, keterlambatan, dan aktivitas bimbingan",
            ),
            (
                "Siswa membuka Dashboard Siswa",
                "Sistem menampilkan ringkasan pinjaman aktif, pengajuan pending, notifikasi, dan akses cepat ajukan alat/bahan",
            ),
        ],
    ),
    (
        "n. Pengujian alpha notifikasi",
        "Pengujian alpha notifikasi",
        [
            (
                "Siswa / Guru / Admin membuka menu Notifikasi",
                "Sistem menampilkan daftar notifikasi milik pengguna tersebut",
            ),
            (
                "Pengguna menandai satu notifikasi sebagai sudah dibaca",
                "Notifikasi berpindah ke status dibaca",
            ),
            (
                "Pengguna menandai semua notifikasi sudah dibaca",
                "Sistem menampilkan pesan “Semua notifikasi ditandai sudah dibaca.”",
            ),
            (
                "Terjadi event pengajuan baru / disetujui / ditolak / terlambat",
                "Penerima terkait menerima notifikasi dengan judul sesuai event",
            ),
        ],
    ),
    (
        "o. Pengujian alpha laporan",
        "Pengujian alpha laporan",
        [
            (
                "Admin membuka menu Laporan",
                "Sistem menampilkan workspace laporan dengan tab Ringkasan, Inventaris, Peminjaman, dan Pengguna",
            ),
            (
                "Admin menerapkan filter tanggal / status / jenis lalu melihat hasil laporan",
                "Sistem menampilkan data laporan sesuai filter",
            ),
            (
                "Admin mengekspor laporan ke PDF / Excel",
                "Sistem mengunduh berkas laporan sesuai format yang dipilih",
            ),
            (
                "Guru membuka menu Laporan",
                "Sistem menampilkan laporan bimbingan / inventaris tanpa tab Pengguna",
            ),
            (
                "Guru mengekspor laporan ke PDF / Excel",
                "Sistem mengunduh berkas laporan sesuai format yang dipilih",
            ),
        ],
    ),
    (
        "p. Pengujian alpha profil pengguna",
        "Pengujian alpha profil",
        [
            (
                "Pengguna membuka halaman Profil",
                "Sistem menampilkan informasi profil akun yang sedang login",
            ),
            (
                "Pengguna mengubah nama / email profil",
                "Sistem menyimpan perubahan profil",
            ),
            (
                "Pengguna mengubah kata sandi dengan password lama benar dan konfirmasi cocok",
                "Sistem memperbarui kata sandi berhasil",
            ),
            (
                "Pengguna mengubah kata sandi dengan konfirmasi tidak cocok",
                "Sistem menolak perubahan dan menampilkan pesan validasi konfirmasi password",
            ),
        ],
    ),
    (
        "q. Pengujian alpha lupa password",
        "Pengujian alpha lupa password",
        [
            (
                "Pengguna membuka tautan Lupa password? dari halaman login",
                "Sistem menampilkan form permintaan tautan reset password",
            ),
            (
                "Pengguna mengirim email yang terdaftar untuk reset password",
                "Sistem menampilkan status bahwa tautan reset telah dikirim (jika mailer aktif)",
            ),
            (
                "Pengguna membuka tautan reset valid lalu mengisi password baru dan konfirmasi",
                "Sistem menyimpan password baru dan mengarahkan ke halaman login",
            ),
        ],
    ),
]


def set_run_font(run, *, bold: bool = False, size: int = 11, name: str = "Times New Roman"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def shade_cell(cell, hex_color: str):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def set_cell_borders(cell):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    existing = tc_pr.find(qn("w:tcBorders"))
    if existing is not None:
        tc_pr.remove(existing)
    tc_borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "000000")
        tc_borders.append(element)
    tc_pr.append(tc_borders)


def set_vertical_merge(cell, restart: bool):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    for child in list(tc_pr):
        if child.tag == qn("w:vMerge"):
            tc_pr.remove(child)
    vmerge = OxmlElement("w:vMerge")
    if restart:
        vmerge.set(qn("w:val"), "restart")
    tc_pr.append(vmerge)


def fill_cell(cell, text: str, *, bold: bool = False, center: bool = False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    )
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(text)
    set_run_font(run, bold=bold, size=11)
    set_cell_borders(cell)


def add_alpha_table(doc: Document, number: int, caption: str, rows: list[tuple[str, str]]):
    # Section caption above is handled by caller; this is table title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(6)
    title.paragraph_format.space_after = Pt(6)
    run = title.add_run(f"Tabel 3.{number} {caption}")
    set_run_font(run, bold=True, size=12)

    # 2 header rows: row0 = Test Case | Yang diharapkan | Hasil Pengujian (merged)
    # row1 = (empty continue) | (empty continue) | Sesuai | Tidak Sesuai
    table = doc.add_table(rows=2 + len(rows), cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    widths = [Cm(5.5), Cm(6.5), Cm(2.2), Cm(2.5)]

    # Header row 0
    h0 = table.rows[0].cells
    fill_cell(h0[0], "Test Case", bold=True, center=True)
    fill_cell(h0[1], "Yang diharapkan", bold=True, center=True)
    fill_cell(h0[2], "Hasil Pengujian", bold=True, center=True)
    fill_cell(h0[3], "", bold=True, center=True)
    h0[2].merge(h0[3])
    for i in range(3):
        shade_cell(h0[i], HEADER_BG)

    # Header row 1 — vertical merge for first two cols, subheaders for result
    h1 = table.rows[1].cells
    fill_cell(h1[0], "", bold=True, center=True)
    fill_cell(h1[1], "", bold=True, center=True)
    fill_cell(h1[2], "Sesuai", bold=True, center=True)
    fill_cell(h1[3], "Tidak Sesuai", bold=True, center=True)
    shade_cell(h1[0], HEADER_BG)
    shade_cell(h1[1], HEADER_BG)
    shade_cell(h1[2], HEADER_BG)
    shade_cell(h1[3], HEADER_BG)

    # Vertical merge header columns 0 and 1 across two header rows
    set_vertical_merge(h0[0], True)
    set_vertical_merge(h1[0], False)
    set_vertical_merge(h0[1], True)
    set_vertical_merge(h1[1], False)

    # Restore visible text in merged header cells (Word sometimes drops)
    fill_cell(h0[0], "Test Case", bold=True, center=True)
    fill_cell(h0[1], "Yang diharapkan", bold=True, center=True)
    shade_cell(h0[0], HEADER_BG)
    shade_cell(h0[1], HEADER_BG)
    shade_cell(h1[0], HEADER_BG)
    shade_cell(h1[1], HEADER_BG)

    for idx, (case, expected) in enumerate(rows):
        cells = table.rows[2 + idx].cells
        fill_cell(cells[0], case)
        fill_cell(cells[1], expected)
        fill_cell(cells[2], "", center=True)
        fill_cell(cells[3], "", center=True)
        for col, width in enumerate(widths):
            cells[col].width = width

    for col, width in enumerate(widths):
        table.rows[0].cells[col].width = width
        table.rows[1].cells[col].width = width


def build_document() -> Document:
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = heading.add_run("3.x Pengujian Alpha (Black-Box Testing)")
    set_run_font(run, bold=True, size=14)

    intro = doc.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    intro_run = intro.add_run(
        "Pengujian alpha dilakukan terhadap Sistem Manajemen Peminjaman Alat dan Bahan "
        "Laboratorium Teknik Audio Video (SMK Negeri 7 Bekasi) dengan metode black-box testing. "
        "Pengujian berfokus pada kesesuaian fungsi sistem terhadap kebutuhan pengguna tanpa "
        "melihat struktur kode program. Skenario mencakup autentikasi, otorisasi role "
        "(Admin, Guru, Siswa), pengelolaan master data, pengajuan dan verifikasi peminjaman, "
        "jaminan kartu, pengembalian & inspeksi, notifikasi, laporan, serta profil pengguna. "
        "Kolom Hasil Pengujian diisi oleh penguji dengan tanda centang pada kolom Sesuai atau Tidak Sesuai."
    )
    set_run_font(intro_run, size=12)

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    note_run = note.add_run(
        "Catatan pengujian: login menggunakan username (bukan email). Role yang diuji adalah "
        "admin, guru, dan siswa. Pesan sukses operasional umumnya berbahasa Indonesia; "
        "pesan autentikasi bawaan Laravel dapat tampil dalam bahasa Inggris sesuai locale aplikasi."
    )
    set_run_font(note_run, size=11)

    start_number = 21
    for offset, (section_label, caption, rows) in enumerate(SECTIONS):
        label_p = doc.add_paragraph()
        label_p.paragraph_format.space_before = Pt(14)
        label_p.paragraph_format.space_after = Pt(2)
        label_run = label_p.add_run(section_label)
        set_run_font(label_run, bold=True, size=12)

        add_alpha_table(doc, start_number + offset, caption, rows)

    # Closing summary table hint
    closing = doc.add_paragraph()
    closing.paragraph_format.space_before = Pt(18)
    closing.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    closing_run = closing.add_run(
        "Setelah seluruh skenario diuji, hasil pengujian alpha dinyatakan berhasil apabila "
        "seluruh test case kritikal pada alur login, RBAC, pengajuan, verifikasi, jaminan, "
        "pengembalian/inspeksi, serta hak akses tiap role bernilai Sesuai. Temuan Tidak Sesuai "
        "dicatat sebagai defect untuk perbaikan sebelum pengujian beta."
    )
    set_run_font(closing_run, size=12)

    return doc


def main():
    out = Path(__file__).resolve().parent / "pengujian-alpha-blackbox.docx"
    doc = build_document()
    doc.save(out)
    print(f"Saved: {out}")
    total_cases = sum(len(rows) for _, _, rows in SECTIONS)
    print(f"Sections: {len(SECTIONS)} | Test cases: {total_cases}")


if __name__ == "__main__":
    main()
